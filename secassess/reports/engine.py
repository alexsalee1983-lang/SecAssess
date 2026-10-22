"""L6 报告引擎：从单一事实源汇总报告模型，输出 Markdown / HTML（可选 DOCX）。

每条发现附证据链：通道 → 引擎 → 时间 → 原始证据 → 资产节点。
"""
from __future__ import annotations

import html
from dataclasses import dataclass, field

from ..core.scope_gate import ScopePolicy
from ..data.asset_graph import AssetGraph
from ..data.evidence import EvidenceStore
from ..models import Category, FindingStatus, NormalizedFinding, Severity, utcnow

SEVERITY_ORDER = [s.value for s in Severity]
_CATEGORY_LABEL = {"asset": "资产", "vulnerability": "漏洞", "misconfig": "配置基线",
                   "exposure": "暴露面", "compliance": "合规"}


@dataclass
class ReportModel:
    project_id: str
    authorization_ref: str
    generated_at: str
    assets: list[dict]
    services: list[dict]
    findings: list[dict]
    severity_counts: dict[str, int] = field(default_factory=dict)
    category_counts: dict[str, int] = field(default_factory=dict)


def evidence_chain(f: NormalizedFinding, graph: AssetGraph) -> list[str]:
    node = graph.get_node(f.asset_ref) if f.asset_ref else None
    chain = []
    for s in f.sources:
        chain.append(f"{s.channel.value} → {s.adapter}/{s.engine}@{s.engine_version}")
    chain.append(f"首见 {f.first_seen.isoformat()} / 最近 {f.last_seen.isoformat()}")
    if f.raw_ref:
        chain.append(f"原始证据 {f.raw_ref}")
    chain.append(f"资产节点 {node['kind']}:{node['key']}" if node else "资产节点 （未关联）")
    return chain


def build_report(project_id: str, evidence: EvidenceStore, graph: AssetGraph,
                 policy: ScopePolicy | None = None, include_closed: bool = False) -> ReportModel:
    findings = [f for f in evidence.query(project_id)
                if include_closed or f.status not in (FindingStatus.FALSE_POSITIVE,)]
    issues = [f for f in findings if f.category != Category.ASSET]
    sev_counts = {s: 0 for s in SEVERITY_ORDER}
    for f in issues:
        sev_counts[f.severity.value] += 1
    cat_counts: dict[str, int] = {}
    for f in findings:
        cat_counts[f.category.value] = cat_counts.get(f.category.value, 0) + 1
    rows = [{**f.to_dict(), "evidence_chain": evidence_chain(f, graph)} for f in issues]
    return ReportModel(
        project_id=project_id, authorization_ref=policy.authorization_ref if policy else "（未录入）",
        generated_at=utcnow().isoformat(), assets=graph.nodes(project_id, "asset"),
        services=graph.nodes(project_id, "service"), findings=rows,
        severity_counts=sev_counts, category_counts=cat_counts)


def _md(text) -> str:
    return str(text if text is not None else "").replace("|", "\\|").replace("\n", " ")


def render_markdown(r: ReportModel) -> str:
    lines = [f"# 安全评估报告：{_md(r.project_id)}", "",
             f"- 授权书编号：{_md(r.authorization_ref)}", f"- 生成时间：{r.generated_at}",
             f"- 资产数：{len(r.assets)}；服务数：{len(r.services)}；问题数：{len(r.findings)}", "",
             "## 概览", "", "| 严重性 | 数量 |", "|---|---|"]
    lines += [f"| {s} | {r.severity_counts.get(s, 0)} |" for s in SEVERITY_ORDER]
    lines += ["", "## 资产清单", "", "| 资产 | 服务 |", "|---|---|"]
    for a in r.assets:
        ports = sorted(s["key"].rsplit(":", 1)[-1] for s in r.services if s["attrs"].get("host") == a["key"])
        lines.append(f"| {_md(a['key'])} | {_md(', '.join(ports))} |")
    lines += ["", "## 发现明细", ""]
    for i, f in enumerate(r.findings, 1):
        lines += [f"### {i}. [{f['severity']}] {_md(f['title'])}", "",
                  f"- 类别：{_CATEGORY_LABEL.get(f['category'], f['category'])}；目标：`{_md(f['target'])}`；"
                  f"置信度：{f['confidence']}；状态：{f['status']}",
                  f"- 整改建议：{_md(f['remediation'] or '—')}", "- 证据链："]
        lines += [f"  - {_md(c)}" for c in f["evidence_chain"]]
        lines.append("")
    return "\n".join(lines)


def render_html(r: ReportModel) -> str:
    e = html.escape
    sev_rows = "".join(f"<tr><td>{s}</td><td>{r.severity_counts.get(s, 0)}</td></tr>" for s in SEVERITY_ORDER)
    asset_rows = "".join(f"<tr><td>{e(a['key'])}</td></tr>" for a in r.assets)
    items = []
    for f in r.findings:
        chain = "".join(f"<li>{e(c)}</li>" for c in f["evidence_chain"])
        items.append(
            f"<section class='finding sev-{e(f['severity'])}'><h3>[{e(f['severity'])}] {e(f['title'])}</h3>"
            f"<p>类别：{e(_CATEGORY_LABEL.get(f['category'], f['category']))}；目标：<code>{e(f['target'])}</code>；"
            f"置信度：{e(f['confidence'])}；状态：{e(f['status'])}</p>"
            f"<p>整改建议：{e(f['remediation'] or '—')}</p><ul class='chain'>{chain}</ul></section>")
    return (
        "<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'>"
        f"<title>安全评估报告 {e(r.project_id)}</title><style>"
        "body{font-family:system-ui,sans-serif;max-width:960px;margin:2em auto;padding:0 1em;color:#222}"
        "table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:4px 10px}"
        ".finding{border-left:4px solid #999;padding:0 1em;margin:1em 0}.sev-critical{border-color:#b00020}"
        ".sev-high{border-color:#e65100}.sev-medium{border-color:#f9a825}.sev-low{border-color:#2e7d32}"
        ".chain{color:#555;font-size:.9em}</style></head><body>"
        f"<h1>安全评估报告：{e(r.project_id)}</h1>"
        f"<p>授权书编号：{e(r.authorization_ref)}<br>生成时间：{e(r.generated_at)}<br>"
        f"资产数：{len(r.assets)}；服务数：{len(r.services)}；问题数：{len(r.findings)}</p>"
        f"<h2>概览</h2><table><tr><th>严重性</th><th>数量</th></tr>{sev_rows}</table>"
        f"<h2>资产清单</h2><table><tr><th>资产</th></tr>{asset_rows}</table>"
        f"<h2>发现明细</h2>{''.join(items)}</body></html>")


def render_docx(r: ReportModel, path: str) -> str:
    """需要可选依赖 python-docx（pip install 'secassess[docx]'）。"""
    import docx  # noqa: PLC0415 - 可选依赖按需导入

    doc = docx.Document()
    doc.add_heading(f"安全评估报告：{r.project_id}", 0)
    doc.add_paragraph(f"授权书编号：{r.authorization_ref}\n生成时间：{r.generated_at}\n"
                      f"资产数：{len(r.assets)}；问题数：{len(r.findings)}")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "严重性", "数量"
    for s in SEVERITY_ORDER:
        row = table.add_row().cells
        row[0].text, row[1].text = s, str(r.severity_counts.get(s, 0))
    doc.add_heading("发现明细", 1)
    for f in r.findings:
        doc.add_heading(f"[{f['severity']}] {f['title']}", 2)
        doc.add_paragraph(f"目标：{f['target']}；置信度：{f['confidence']}；状态：{f['status']}")
        doc.add_paragraph(f"整改建议：{f['remediation'] or '—'}")
        for c in f["evidence_chain"]:
            doc.add_paragraph(c, style="List Bullet")
    doc.save(path)
    return path
